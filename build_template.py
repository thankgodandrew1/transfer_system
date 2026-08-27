"""One-off script to (re)generate templates/template.docx with the current
placeholder layout: zone heading paragraph + a one-row placeholder table,
no Status column. Run manually whenever the placeholder set changes.
"""
from pathlib import Path

from docx import Document

document = Document()
zone_paragraph = document.add_paragraph("{{ZONE}}")

table = document.add_table(rows=2, cols=5)
table.style = "Table Grid"
headers = [
    "NAME OF MISSIONARY",
    "ASSIGNMENT",
    "NEW/EXISTING ZONE",
    "NEW/EXISTING AREA",
    "NEW/EXISTING COMPANION(S)",
]
for index, header in enumerate(headers):
    table.rows[0].cells[index].text = header

placeholders = [
    "{{MISSIONARY_NAME}}",
    "{{ASSIGNMENT}}",
    "{{ZONE}}",
    "{{AREA}}",
    "{{COMPANION}}",
]
for index, placeholder in enumerate(placeholders):
    table.rows[1].cells[index].text = placeholder

output_path = Path("templates/template.docx")
document.save(output_path)
print(f"Wrote {output_path}")
