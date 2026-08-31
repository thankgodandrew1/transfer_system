from __future__ import annotations

import csv
import difflib
import html
import io
import json
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Sequence

import openpyxl
import pdfplumber
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


GREEN = "1A5C38"
GOLD = "D4AF37"
INK = "1A1A1A"
MUTED = "555555"
LIGHT_MUTED = "888888"
ZEBRA = "EEF4F0"
BORDER = "CCCCCC"

DIRECTORY_COLUMNS = ("Zone", "Apartment", "Area")
DEFAULT_STATUS = ""


@dataclass(frozen=True)
class Assignment:
    name: str
    zone: str
    area: str
    order: int

    @property
    def key(self) -> str:
        return canonical(self.name)


@dataclass(frozen=True)
class DirectoryEntry:
    zone: str
    apartment: str
    area: str


@dataclass(frozen=True)
class ManualOverride:
    name: str
    from_apartment: str
    to_apartment: str
    status: str = ""

    @property
    def key(self) -> str:
        return canonical(self.name)


@dataclass(frozen=True)
class Movement:
    missionary: str
    previous_zone: str
    from_apartment: str
    to_apartment: str
    status: str
    order: int
    manual: bool = False


@dataclass
class ParseResult:
    assignments: list[Assignment]
    title: str = ""
    expected_count: int | None = None
    warnings: list[str] = field(default_factory=list)


@dataclass
class MovementComparison:
    movements: list[Movement]
    previous_total: int
    current_total: int
    matched_total: int
    released_total: int
    new_arrivals_total: int
    unchanged_total: int
    warnings: list[str] = field(default_factory=list)

    @property
    def zone_count(self) -> int:
        return len({canonical(item.previous_zone) for item in self.movements})


class ApartmentDirectory:
    def __init__(self, entries: Iterable[DirectoryEntry]) -> None:
        self.entries: list[DirectoryEntry] = []
        self.by_area: dict[str, DirectoryEntry] = {}
        self.by_fragment: dict[str, DirectoryEntry] = {}
        for entry in entries:
            cleaned = DirectoryEntry(
                zone=clean_text(entry.zone),
                apartment=clean_text(entry.apartment),
                area=clean_text(entry.area),
            )
            if not all((cleaned.zone, cleaned.apartment, cleaned.area)):
                continue
            key = canonical(cleaned.area)
            existing = self.by_area.get(key)
            if existing and canonical(existing.apartment) != canonical(cleaned.apartment):
                raise ValueError(
                    f'Apartment directory assigns area "{cleaned.area}" to both '
                    f'"{existing.apartment}" and "{cleaned.apartment}".'
                )
            if not existing:
                self.entries.append(cleaned)
                self.by_area[key] = cleaned
                for fragment in re.split(r"[/;]", cleaned.area):
                    fragment_key = canonical(fragment)
                    if fragment_key:
                        self.by_fragment.setdefault(fragment_key, cleaned)
        if not self.entries:
            raise ValueError("The apartment directory contains no usable Zone, Apartment, Area rows.")

    def find(self, area: str) -> DirectoryEntry | None:
        key = canonical(area)
        if not key:
            return None
        replacements = {
            "UDOMANA": "UDOUMANA",
            "IBIAKUURUAN": "IBIAKUARAN",
            "ONOIBIONO": "ONAIBIONO",
        }
        variants = [key]
        for old, new in replacements.items():
            variants.extend(item.replace(old, new) for item in list(variants) if old in item)
        for item in list(variants):
            if item.endswith("ROAD"):
                variants.append(item[:-4])
            if item.endswith("UYO"):
                variants.append(item[:-3])
            variants.append(re.sub(r"(?<![12])\d+$", "", item))
        variants = list(OrderedDict.fromkeys(item for item in variants if item))
        for item in variants:
            if item in self.by_area:
                return self.by_area[item]
            if item in self.by_fragment:
                return self.by_fragment[item]

        # Combined historical labels often list two areas separated by a slash.
        # Prefer the earliest complete directory label/fragment in the cell; this
        # represents the primary area printed first in the Transfer News.
        contained: list[tuple[int, int, DirectoryEntry]] = []
        for item in variants:
            for candidate, entry in {**self.by_area, **self.by_fragment}.items():
                if len(candidate) >= 7 and candidate in item:
                    contained.append((item.index(candidate), -len(candidate), entry))
        if contained:
            return min(contained, key=lambda value: (value[0], value[1]))[2]

        # Tolerate a small spelling difference only when the best match is both
        # strong and clearly better than the runner-up.
        scored = sorted(
            (
                (difflib.SequenceMatcher(None, item, candidate).ratio(), candidate, entry)
                for item in variants
                for candidate, entry in self.by_area.items()
            ),
            key=lambda value: value[0],
            reverse=True,
        )
        if scored and scored[0][0] >= 0.84:
            runner_up = scored[1][0] if len(scored) > 1 else 0
            if scored[0][0] - runner_up >= 0.025:
                return scored[0][2]
        return None

    def to_csv_bytes(self) -> bytes:
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(DIRECTORY_COLUMNS)
        for item in self.entries:
            writer.writerow((item.zone, item.apartment, item.area))
        return stream.getvalue().encode("utf-8-sig")

    def display_zone(self, zone: str) -> str:
        raw = clean_text(zone)
        key = canonical(raw)
        key = key.replace("UDOMANA", "UDOUMANA")
        for entry in self.entries:
            candidate = canonical(entry.zone)
            candidate_base = re.sub(r"ZONE$", "", candidate)
            if key in {candidate, candidate_base}:
                return entry.zone
        return raw


def clean_text(value: object) -> str:
    text = "" if value is None else str(value)
    for _ in range(2):
        text = html.unescape(text)
    text = text.replace("\u00a0", " ").replace("\n", " ").replace("\r", " ")
    text = re.sub(r"\s+", " ", text).strip(" \t|;,")
    return text


def canonical(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", clean_text(value).upper())


def _title_case_code(value: str) -> str:
    words = clean_text(value).upper().split()
    return " ".join(word if len(word) <= 3 and word in {"1", "2", "3", "4"} else word.title() for word in words)


def _deduplicate_assignments(assignments: Iterable[Assignment]) -> list[Assignment]:
    by_name: OrderedDict[str, Assignment] = OrderedDict()
    for assignment in assignments:
        if not assignment.key:
            continue
        if assignment.key in {"ELDER", "SISTER"}:
            raise ValueError(
                "A Transfer News row contains only the title "
                f'"{assignment.name}" without the missionary surname. '
                "This usually means a table row was split or clipped during PDF export."
            )
        existing = by_name.get(assignment.key)
        if existing and (canonical(existing.zone), canonical(existing.area)) != (
            canonical(assignment.zone),
            canonical(assignment.area),
        ):
            raise ValueError(
                f'"{assignment.name}" appears more than once with different zone or area values.'
            )
        by_name.setdefault(assignment.key, assignment)
    return list(by_name.values())


def _header_indexes(row: Sequence[object]) -> tuple[int, int, int] | None:
    cells = [canonical(cell) for cell in row]
    missionary = next((index for index, cell in enumerate(cells) if "MISSIONARY" in cell), None)
    zone = next((index for index, cell in enumerate(cells) if cell == "ZONE" or cell.endswith("ZONE")), None)
    area = next((index for index, cell in enumerate(cells) if cell == "AREA" or cell.endswith("AREA")), None)
    if missionary is None or zone is None or area is None:
        return None
    return missionary, zone, area


def _extract_transfer_title(text: str) -> str:
    compact = clean_text(text).upper()
    months = (
        "JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|"
        "OCTOBER|NOVEMBER|DECEMBER"
    )
    patterns = (
        rf"((?:{months})(?:\s*(?:/|&|AND|-)\s*(?:{months}))?\s+20\d{{2}}\s+TRANSFER\s+NEWS)",
        rf"(TRANSFER\s+NEWS\s*[·\-:]?\s*(?:{months})(?:\s*(?:/|&|AND|-)\s*(?:{months}))?\s+20\d{{2}})",
    )
    for pattern in patterns:
        match = re.search(pattern, compact)
        if match:
            return clean_text(match.group(1))
    return ""


def _extract_expected_count(text: str) -> int | None:
    patterns = (
        r"TOTAL\s+MISSIONAR(?:Y|IES)\s*[:\-]?\s*(\d{2,3})",
        r"(\d{2,3})\s+TOTAL\s+MISSIONAR(?:Y|IES)",
        r"MISSIONARY\s+(?:COMPLEMENT|STRENGTH)\s*[:\-]?\s*(\d{2,3})",
    )
    upper = clean_text(text).upper()
    for pattern in patterns:
        match = re.search(pattern, upper)
        if match:
            return int(match.group(1))
    return None


def parse_transfer_pdf(path: str | Path) -> ParseResult:
    source = Path(path)
    assignments: list[Assignment] = []
    first_pages: list[str] = []
    order = 0
    with pdfplumber.open(source) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            page_text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            if page_number <= 2:
                first_pages.append(page_text)
            for table in page.extract_tables() or []:
                indexes: tuple[int, int, int] | None = None
                for row in table:
                    if not row:
                        continue
                    detected = _header_indexes(row)
                    if detected:
                        indexes = detected
                        continue
                    if indexes is None:
                        continue
                    name_index, zone_index, area_index = indexes
                    if max(indexes) >= len(row):
                        continue
                    sequence = clean_text(row[0])
                    sequence_digits = re.sub(r"\s+", "", sequence)
                    if not sequence_digits and assignments:
                        # WPS/Word can split the final table row across two PDF
                        # pages. The continuation begins with a blank sequence
                        # cell and carries the remaining surname/zone fragment.
                        name_fragment = clean_text(row[name_index])
                        zone_fragment = clean_text(row[zone_index])
                        area_fragment = clean_text(row[area_index])
                        if name_fragment or zone_fragment or area_fragment:
                            previous = assignments[-1]
                            assignments[-1] = Assignment(
                                name=clean_text(f"{previous.name} {name_fragment}"),
                                zone=clean_text(f"{previous.zone} {zone_fragment}"),
                                area=clean_text(f"{previous.area} {area_fragment}"),
                                order=previous.order,
                            )
                        continue
                    if not re.fullmatch(r"\d{1,3}", sequence_digits):
                        continue
                    name = clean_text(row[name_index])
                    zone = clean_text(row[zone_index])
                    area = clean_text(row[area_index])
                    if not name or not zone or not area:
                        continue
                    order += 1
                    assignments.append(Assignment(name=name, zone=zone, area=area, order=order))

    incomplete = [item for item in assignments if item.key in {"ELDER", "SISTER"}]
    if incomplete:
        titles = ", ".join(item.name.upper() for item in incomplete[:5])
        raise ValueError(
            f'{source.name} contains {len(incomplete)} incomplete missionary name row(s) ({titles}). '
            "The surname is missing from the PDF table. Re-export the Transfer News as PDF, or paste a "
            "complete Missionary | Zone | Area table in the fallback field."
        )
    assignments = _deduplicate_assignments(assignments)
    cover_text = "\n".join(first_pages)
    expected = _extract_expected_count(cover_text)
    warnings: list[str] = []
    if not assignments:
        raise ValueError(
            f'No missionary assignment table could be read from "{source.name}". '
            "Paste that Transfer News as text in the fallback field and try again."
        )
    if expected and len(assignments) != expected:
        warnings.append(
            f'{source.name}: the cover reports {expected} missionaries, while {len(assignments)} named table rows were read. '
            "New-arrival placeholders without names are not compared."
        )
    return ParseResult(
        assignments=assignments,
        title=_extract_transfer_title(cover_text),
        expected_count=expected,
        warnings=warnings,
    )


def _parse_delimited_transfer_text(text: str) -> list[Assignment]:
    rows = [line for line in text.splitlines() if clean_text(line)]
    if not rows:
        return []
    delimiter = "\t" if any("\t" in row for row in rows[:5]) else "|"
    if delimiter not in "\n".join(rows[:5]):
        delimiter = ","
    parsed = list(csv.reader(rows, delimiter=delimiter))
    header_at = None
    indexes = None
    for index, row in enumerate(parsed[:12]):
        detected = _header_indexes(row)
        if detected:
            header_at, indexes = index, detected
            break
    if indexes is None:
        return []
    result: list[Assignment] = []
    for row in parsed[(header_at or 0) + 1 :]:
        if max(indexes) >= len(row):
            continue
        name, zone, area = (clean_text(row[index]) for index in indexes)
        if name and zone and area:
            result.append(Assignment(name=name, zone=zone, area=area, order=len(result) + 1))
    return result


def _parse_pdf_style_text(text: str, directory: ApartmentDirectory | None) -> list[Assignment]:
    if directory is None:
        return []
    upper = text.upper()
    lines = [clean_text(line) for line in upper.splitlines() if clean_text(line)]
    known_areas = sorted(directory.entries, key=lambda item: len(item.area), reverse=True)
    known_zones = sorted({item.zone for item in directory.entries}, key=len, reverse=True)
    assignments: list[Assignment] = []
    for line in lines:
        line = re.sub(r"^\s*\d{1,3}\s+", "", line)
        area_entry = next((entry for entry in known_areas if canonical(entry.area) in canonical(line)), None)
        if area_entry is None:
            continue
        role_match = re.search(r"\b(?:ELDER|SISTER)\b", line)
        if not role_match:
            continue
        zone = next((item for item in known_zones if canonical(item) in canonical(line)), area_entry.zone)
        area_position = canonical(line).find(canonical(area_entry.area))
        compact_line = canonical(line)
        name_compact = compact_line[:area_position]
        name_compact = re.sub(r"^\d+", "", name_compact)
        # Stop the name at the first assignment/role code or previous-zone marker when available.
        name_compact = re.split(
            r"(?:AP|DL|DT|STL|ZL|ZLT|TRAINER|SENIOR|JUNIOR|COMPANION|" +
            "|".join(re.escape(canonical(zone_name)) for zone_name in known_zones) + r")",
            name_compact,
            maxsplit=1,
        )[0]
        if not (name_compact.startswith("ELDER") or name_compact.startswith("SISTER")):
            continue
        honorific = "ELDER" if name_compact.startswith("ELDER") else "SISTER"
        remainder = name_compact[len(honorific) :]
        if len(remainder) < 2:
            continue
        display_name = f"{honorific.title()} {remainder.title()}"
        assignments.append(
            Assignment(display_name, zone, area_entry.area, len(assignments) + 1)
        )
    return _deduplicate_assignments(assignments)


def parse_transfer_text(text: str, directory: ApartmentDirectory | None = None) -> ParseResult:
    cleaned = text.strip()
    if not cleaned:
        raise ValueError("The plain-text Transfer News fallback is empty.")
    assignments = _parse_delimited_transfer_text(cleaned)
    if not assignments:
        assignments = _parse_pdf_style_text(cleaned, directory)
    assignments = _deduplicate_assignments(assignments)
    if not assignments:
        raise ValueError(
            "The pasted Transfer News could not be read. Use a table with Missionary, Zone, and Area columns "
            "separated by tabs, commas, or vertical bars."
        )
    return ParseResult(assignments=assignments, title=_extract_transfer_title(cleaned))


def _entries_from_csv(path: Path) -> list[DirectoryEntry]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        names = {canonical(name): name for name in (reader.fieldnames or [])}
        required = {canonical(name) for name in DIRECTORY_COLUMNS}
        if not required.issubset(names):
            raise ValueError("Apartment CSV must contain Zone, Apartment, and Area columns.")
        return [
            DirectoryEntry(
                zone=row.get(names[canonical("Zone")], ""),
                apartment=row.get(names[canonical("Apartment")], ""),
                area=row.get(names[canonical("Area")], ""),
            )
            for row in reader
        ]


def _entries_from_xlsx(path: Path) -> list[DirectoryEntry]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            names = {canonical(value): index for index, value in enumerate(header) if clean_text(value)}
            if not {canonical(name) for name in DIRECTORY_COLUMNS}.issubset(names):
                continue
            result: list[DirectoryEntry] = []
            for row in rows:
                result.append(
                    DirectoryEntry(
                        clean_text(row[names[canonical("Zone")]]) if len(row) > names[canonical("Zone")] else "",
                        clean_text(row[names[canonical("Apartment")]]) if len(row) > names[canonical("Apartment")] else "",
                        clean_text(row[names[canonical("Area")]]) if len(row) > names[canonical("Area")] else "",
                    )
                )
            return result
    finally:
        workbook.close()
    raise ValueError("No worksheet contains Zone, Apartment, and Area columns.")


def _entries_from_json(path: Path) -> list[DirectoryEntry]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError("The apartment JSON file could not be read.") from exc
    result: list[DirectoryEntry] = []
    lookup = payload.get("area_to_apartment") if isinstance(payload, dict) else None
    if isinstance(lookup, dict):
        for area, value in lookup.items():
            if isinstance(value, dict):
                result.append(DirectoryEntry(value.get("zone", ""), value.get("apartment", ""), area))
            elif isinstance(value, str):
                result.append(DirectoryEntry("Unspecified Zone", value, area))
    if not result and isinstance(payload, dict) and isinstance(payload.get("zones"), dict):
        for zone, zone_value in payload["zones"].items():
            for apartment in (zone_value or {}).get("apartments", []):
                for area in apartment.get("areas", []):
                    result.append(DirectoryEntry(zone, apartment.get("name", ""), area))
    return result


def _entries_from_docx(path: Path) -> list[DirectoryEntry]:
    document = Document(path)
    result: list[DirectoryEntry] = []
    current_zone = ""
    current_apartment = ""
    in_note = False
    paragraphs = document.paragraphs
    for paragraph_index, paragraph in enumerate(paragraphs):
        text = clean_text(paragraph.text)
        if not text:
            in_note = False
            continue
        upper = text.upper().rstrip(":")
        if upper.startswith("NOTE:"):
            in_note = True
            continue
        if in_note:
            continue
        max_size = max(
            (run.font.size.pt for run in paragraph.runs if run.font.size is not None),
            default=0,
        )
        if "ZONE" in upper and len(upper.split()) <= 5:
            current_zone = upper
            current_apartment = ""
            continue
        next_text = ""
        for following in paragraphs[paragraph_index + 1 :]:
            next_text = clean_text(following.text).upper()
            if next_text:
                break
        if max_size >= 20 and "APARTMENT" not in upper and "APARTMENT" in next_text:
            current_zone = upper + ("" if upper.endswith("ZONE") else " ZONE")
            current_apartment = ""
            continue
        if "APARTMENT" in upper:
            current_apartment = re.sub(r"\bAPARTMENT\b", "", upper).strip(" :-")
            current_apartment = re.sub(r"^(?:IM|IIM)\s+", "", current_apartment)
            continue
        if current_zone and max_size >= 18 and len(upper.split()) <= 5:
            current_apartment = upper
            continue
        if current_zone and current_apartment:
            for area in re.split(r"\s*(?:,|;|\u2022|\|)\s*", upper):
                area = clean_text(re.sub(r"^\d+[.)-]?\s*", "", area))
                if area:
                    result.append(DirectoryEntry(current_zone, current_apartment, area))
    return result


def load_apartment_directory(path: str | Path) -> ApartmentDirectory:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".csv":
        entries = _entries_from_csv(source)
    elif suffix == ".xlsx":
        entries = _entries_from_xlsx(source)
    elif suffix == ".json":
        entries = _entries_from_json(source)
    elif suffix == ".docx":
        entries = _entries_from_docx(source)
    else:
        raise ValueError("Apartment directory must be a CSV, XLSX, JSON, or DOCX file.")
    return ApartmentDirectory(entries)


def parse_manual_overrides(text: str, default_status: str = DEFAULT_STATUS) -> list[ManualOverride]:
    overrides: list[ManualOverride] = []
    seen: set[str] = set()
    fallback_status = clean_text(default_status).upper()
    for line_number, original in enumerate(text.splitlines(), start=1):
        line = clean_text(original)
        if not line or line.startswith("#"):
            continue
        status = fallback_status
        parts = [clean_text(part) for part in line.split("|")]
        if len(parts) in {3, 4} and not any(symbol in line for symbol in ("->", "→", "⇒")):
            name, from_apartment, to_apartment = parts[:3]
            if len(parts) == 4:
                status = parts[3].upper() or fallback_status
        else:
            if len(parts) == 2:
                line, status_value = parts
                status = status_value.upper() or fallback_status
            match = re.match(r"^(.+?)\s*:\s*(.+?)\s*(?:->|→|⇒)\s*(.+?)$", line)
            if not match:
                raise ValueError(
                    f"Manual override line {line_number} is invalid. Use Name: FROM -> TO | STATUS."
                )
            name, from_apartment, to_apartment = (clean_text(value) for value in match.groups())
        if not all((name, from_apartment, to_apartment)):
            raise ValueError(f"Manual override line {line_number} has an empty field.")
        item = ManualOverride(name, from_apartment, to_apartment, status[:40])
        if item.key in seen:
            raise ValueError(f'Manual override for "{name}" appears more than once.')
        seen.add(item.key)
        overrides.append(item)
    return overrides


def compare_assignments(
    previous: Sequence[Assignment],
    current: Sequence[Assignment],
    directory: ApartmentDirectory,
    overrides: Sequence[ManualOverride] = (),
    default_status: str = DEFAULT_STATUS,
) -> MovementComparison:
    previous_by_name = {item.key: item for item in _deduplicate_assignments(previous)}
    current_by_name = {item.key: item for item in _deduplicate_assignments(current)}
    overrides_by_name = {item.key: item for item in overrides}
    unknown_overrides = [item.name for item in overrides if item.key not in previous_by_name or item.key not in current_by_name]
    if unknown_overrides:
        raise ValueError(
            "Manual overrides must name missionaries present in both Transfer News files. Not matched: "
            + ", ".join(unknown_overrides)
            + "."
        )

    matched_keys = previous_by_name.keys() & current_by_name.keys()
    movements: list[Movement] = []
    missing: list[str] = []
    unchanged = 0
    # Status cells are intentionally left blank for drivers to mark only after
    # each missionary has moved successfully.
    fallback_status = ""
    for key, old in previous_by_name.items():
        new = current_by_name.get(key)
        if new is None:
            continue
        override = overrides_by_name.get(key)
        old_entry = directory.find(old.area)
        new_entry = directory.find(new.area)
        if override:
            old_zone = directory.display_zone(old.zone)
            movements.append(
                Movement(
                    missionary=new.name,
                    previous_zone=old_zone,
                    from_apartment=override.from_apartment,
                    to_apartment=override.to_apartment,
                    status=fallback_status,
                    order=old.order,
                    manual=True,
                )
            )
            continue
        if canonical(old.area) == canonical(new.area):
            unchanged += 1
            continue
        if old_entry is None:
            missing.append(f'previous area "{old.area}" ({old.name})')
        if new_entry is None:
            missing.append(f'current area "{new.area}" ({new.name})')
        if old_entry is None or new_entry is None:
            continue
        if canonical(old_entry.apartment) == canonical(new_entry.apartment):
            unchanged += 1
            continue
        movements.append(
                Movement(
                    missionary=new.name,
                    previous_zone=directory.display_zone(old.zone),
                from_apartment=old_entry.apartment,
                to_apartment=new_entry.apartment,
                status=fallback_status,
                order=old.order,
            )
        )

    if missing:
        distinct = list(OrderedDict.fromkeys(missing))
        preview = "; ".join(distinct[:12])
        extra = f"; and {len(distinct) - 12} more" if len(distinct) > 12 else ""
        raise ValueError(
            "The apartment directory is missing transfer areas: " + preview + extra + ". "
            "Add those rows to the directory in Excel, then upload it again."
        )

    previous_keys = set(previous_by_name)
    current_keys = set(current_by_name)
    return MovementComparison(
        movements=movements,
        previous_total=len(previous_by_name),
        current_total=len(current_by_name),
        matched_total=len(matched_keys),
        released_total=len(previous_keys - current_keys),
        new_arrivals_total=len(current_keys - previous_keys),
        unchanged_total=unchanged,
    )


def group_movements(movements: Sequence[Movement]) -> list[tuple[str, list[Movement]]]:
    groups: OrderedDict[str, list[Movement]] = OrderedDict()
    labels: dict[str, str] = {}
    for movement in sorted(movements, key=lambda item: item.order):
        key = canonical(movement.previous_zone)
        labels.setdefault(key, clean_text(movement.previous_zone).upper())
        groups.setdefault(key, []).append(movement)
    return [(labels[key], items) for key, items in groups.items()]


def _set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_margins(cell, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_borders(cell, *, color: str = BORDER, size: int = 2, edges: Sequence[str] = ("top", "left", "bottom", "right")) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in edges:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def _remove_cell_borders(cell) -> None:
    _set_cell_borders(cell, color="FFFFFF", size=0)


def _set_table_widths(table, widths: Sequence[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")


def _set_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    properties.append(repeat)


def _keep_row_together(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def _set_repeat_table_header(row) -> None:
    _set_repeat_header(row)


def _set_run(run, size: float, color: str = INK, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def _paragraph(
    container,
    text: str = "",
    *,
    size: float = 9.5,
    color: str = INK,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    before: float = 0,
    after: float = 0,
    line: float = 1,
):
    paragraph = container.add_paragraph() if hasattr(container, "add_paragraph") else container
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if text:
        _set_run(paragraph.add_run(text), size, color, bold)
    return paragraph


def _clear_cell(cell) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)


def _add_rule(paragraph, color: str, size: int = 8, edge: str = "bottom") -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "3")
    border.set(qn("w:color"), color)
    borders.append(border)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))
    _set_run(run, 8, MUTED)


def _configure_page(document: Document, mission_name: str, movement_title: str, president: str) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7), WD_TAB_ALIGNMENT.RIGHT)
    _set_run(p.add_run(mission_name.upper()), 7.5, GREEN, True)
    _set_run(p.add_run("\t" + movement_title.upper()), 7.5, MUTED, True)
    _add_rule(p, GOLD, 8)

    footer = section.footer
    p = footer.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.5), WD_TAB_ALIGNMENT.CENTER)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7), WD_TAB_ALIGNMENT.RIGHT)
    _add_rule(p, BORDER, 4, "top")
    _set_run(p.add_run("CONFIDENTIAL · MISSION OFFICE USE"), 7, LIGHT_MUTED, True)
    _set_run(p.add_run("\t" + president.upper() + "\tPAGE "), 7, LIGHT_MUTED)
    _add_page_field(p)


def _cover_banner(document: Document, mission_name: str) -> None:
    table = document.add_table(rows=1, cols=1)
    _set_table_widths(table, (9360,))
    cell = table.cell(0, 0)
    _clear_cell(cell)
    _set_cell_fill(cell, GREEN)
    _remove_cell_borders(cell)
    _set_cell_margins(cell, 200, 240, 200, 240)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run(p.add_run(mission_name.upper()), 14, GOLD, True)
    p = _paragraph(cell, "THE CHURCH OF JESUS CHRIST OF LATTER-DAY SAINTS", size=8.5, color="FFFFFF", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=3)
    p.paragraph_format.keep_with_next = True


def _metric_strip(document: Document, total: int, zone_count: int, effective_label: str) -> None:
    table = document.add_table(rows=1, cols=3)
    _set_table_widths(table, (3120, 3120, 3120))
    values = ((str(total), "MOVEMENTS"), (str(zone_count), "ZONES COVERED"), (effective_label, "EFFECTIVE"))
    fills = (GREEN, GOLD, GREEN)
    for cell, (value, label), fill in zip(table.rows[0].cells, values, fills):
        _clear_cell(cell)
        _set_cell_fill(cell, fill)
        _remove_cell_borders(cell)
        _set_cell_margins(cell, 180, 120, 170, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        foreground = GREEN if fill == GOLD else "FFFFFF"
        _set_run(p.add_run(value), 17, foreground, True)
        p = _paragraph(cell, label, size=7.5, color=foreground, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=2)


def _publication_details(document: Document, prepared_by: str, president: str, document_date: str) -> None:
    table = document.add_table(rows=1, cols=3)
    _set_table_widths(table, (3120, 3120, 3120))
    values = (("PREPARED BY", prepared_by), ("APPROVED BY", president), ("DOCUMENT DATE", document_date))
    for cell, (label, value) in zip(table.rows[0].cells, values):
        _clear_cell(cell)
        _set_cell_borders(cell, color=GOLD, size=8, edges=("top",))
        _set_cell_margins(cell, 150, 90, 90, 90)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run(p.add_run(label), 7.5, GREEN, True)
        _paragraph(cell, value, size=9, color=INK, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=3)


def _special_note(document: Document, overrides: Sequence[ManualOverride]) -> None:
    if not overrides:
        return
    table = document.add_table(rows=1, cols=1)
    _set_table_widths(table, (9360,))
    cell = table.cell(0, 0)
    _clear_cell(cell)
    _set_cell_fill(cell, ZEBRA)
    _set_cell_borders(cell, color=GOLD, size=8, edges=("top",))
    _set_cell_borders(cell, color=GREEN, size=16, edges=("left",))
    _set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    _set_run(p.add_run("SPECIAL MOVEMENT NOTE"), 8.5, GREEN, True)
    descriptions = [f"{item.name}: {item.from_apartment} → {item.to_apartment}" for item in overrides]
    _paragraph(cell, "; ".join(descriptions), size=9, color=MUTED, before=3)


def _zone_band(document: Document, index: int, zone: str, count: int) -> None:
    table = document.add_table(rows=1, cols=1)
    _set_table_widths(table, (9360,))
    cell = table.cell(0, 0)
    _clear_cell(cell)
    _set_cell_fill(cell, GOLD)
    _remove_cell_borders(cell)
    _set_cell_margins(cell, 125, 165, 115, 165)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.keep_with_next = True
    _set_run(paragraph.add_run(f"{index:02d}  {zone.upper()}  "), 12, GREEN, True)
    label = "MOVEMENT" if count == 1 else "MOVEMENTS"
    _set_run(paragraph.add_run(f"{count} {label}"), 9.5, INK, True)


def _movement_table(document: Document, movements: Sequence[Movement]) -> None:
    table = document.add_table(rows=1, cols=5)
    widths = (520, 2400, 2240, 2240, 1960)
    _set_table_widths(table, widths)
    headers = ("#", "MISSIONARY", "FROM APARTMENT", "TO APARTMENT", "STATUS")
    header_row = table.rows[0]
    _set_repeat_table_header(header_row)
    for index, (cell, value) in enumerate(zip(header_row.cells, headers)):
        _clear_cell(cell)
        _set_cell_fill(cell, GREEN)
        _set_cell_borders(cell, color=GOLD, size=3)
        _set_cell_margins(cell, 110, 130, 110, 130)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        _set_run(p.add_run(value), 9.5, GOLD if index == 0 else "FFFFFF", True)
    for sequence, movement in enumerate(movements, start=1):
        row = table.add_row()
        _keep_row_together(row)
        values = (
            str(sequence),
            movement.missionary,
            movement.from_apartment,
            movement.to_apartment,
            movement.status,
        )
        for index, (cell, value) in enumerate(zip(row.cells, values)):
            _clear_cell(cell)
            if sequence % 2 == 0:
                _set_cell_fill(cell, ZEBRA)
            _set_cell_borders(cell, color=BORDER, size=2)
            _set_cell_margins(cell, 90, 130, 90, 130)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in {0, 4} else WD_ALIGN_PARAGRAPH.LEFT
            if sequence < len(movements):
                p.paragraph_format.keep_with_next = True
            color = LIGHT_MUTED if index == 0 else (GREEN if index == 4 else INK)
            _set_run(p.add_run(value), 9.5, color, bold=index in {1, 4})


def build_movement_docx(
    movements: Sequence[Movement],
    output_path: str | Path,
    *,
    mission_name: str,
    previous_title: str,
    current_title: str,
    president: str,
    prepared_by: str,
    effective_date: str | date | None = None,
    overrides: Sequence[ManualOverride] = (),
) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    groups = group_movements(movements)
    if isinstance(effective_date, date):
        effective_display = effective_date.strftime("%d %B %Y").upper()
        effective_metric = effective_date.strftime("%d %b").upper()
    else:
        effective_display = clean_text(effective_date) or "TO BE CONFIRMED"
        effective_metric = effective_display[:14].upper()
    current_label = clean_text(current_title) or datetime.now().strftime("%B %Y TRANSFER NEWS").upper()
    previous_label = clean_text(previous_title) or "PREVIOUS TRANSFER NEWS"
    movement_title = re.sub(r"\s*TRANSFER\s+NEWS\s*$", "", current_label, flags=re.I).strip()
    movement_title = f"TRANSFER MOVEMENT · {movement_title}"

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    _configure_page(document, mission_name, movement_title, president)

    _cover_banner(document, mission_name)
    _paragraph(document, "TRANSFER MOVEMENT", size=30, color=GREEN, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=32, after=3)
    _paragraph(document, movement_title.replace("TRANSFER MOVEMENT · ", ""), size=20, color=GOLD, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    _paragraph(
        document,
        "Mission-wide apartment movement plan derived from the previous and current Transfer News",
        size=10.5,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=28,
    )
    _metric_strip(document, len(movements), len(groups), effective_metric)
    _paragraph(document, "", after=21)
    _publication_details(document, prepared_by, president, datetime.now().strftime("%d %B %Y").upper())
    _paragraph(
        document,
        "Prepared for coordinated mission-office transport and housing movement. Handle as confidential missionary information.",
        size=8.5,
        color=LIGHT_MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=28,
    )

    document.add_page_break()
    _paragraph(document, "ZONE-BY-ZONE MOVEMENT LISTING", size=16, color=GREEN, bold=True, after=4)
    _paragraph(
        document,
        f"Apartment changes from {previous_label} to {current_label}. Listed under the zone each missionary is moving FROM.",
        size=10.5,
        color=MUTED,
        after=13,
    )
    _special_note(document, overrides)
    if overrides:
        _paragraph(document, "", after=8)

    if not groups:
        table = document.add_table(rows=1, cols=1)
        _set_table_widths(table, (9360,))
        cell = table.cell(0, 0)
        _clear_cell(cell)
        _set_cell_fill(cell, ZEBRA)
        _set_cell_borders(cell, color=GREEN, size=8, edges=("left",))
        _set_cell_margins(cell, 220, 220, 220, 220)
        _set_run(cell.paragraphs[0].add_run("NO APARTMENT MOVEMENTS DETECTED"), 11, GREEN, True)
        _paragraph(cell, "All missionaries present in both files remain in the same mapped apartment.", size=9.5, color=MUTED, before=5)
    else:
        for index, (zone, zone_movements) in enumerate(groups, start=1):
            _zone_band(document, index, zone, len(zone_movements))
            _movement_table(document, zone_movements)
            _paragraph(document, "", after=8)

    _paragraph(
        document,
        f"MOVEMENT SUMMARY · {len(movements)} apartment changes across {len(groups)} previous zones · Effective {effective_display}",
        size=8.5,
        color=GREEN,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=6,
    )
    document.core_properties.title = movement_title
    document.core_properties.author = prepared_by
    document.core_properties.subject = "Mission transfer apartment movement plan"
    document.save(output)
    return output


def write_movement_csv(movements: Sequence[Movement], output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(("Previous Zone", "Missionary", "From Apartment", "To Apartment", "Status", "Manual Override"))
        for item in movements:
            writer.writerow(
                (
                    item.previous_zone,
                    item.missionary,
                    item.from_apartment,
                    item.to_apartment,
                    item.status,
                    "Yes" if item.manual else "No",
                )
            )
    return output
