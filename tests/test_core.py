from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from pypdf import PdfReader

import generate_news_format
from generate_transfer_document import build_transfer_docx, compute_stats
from generate_transfer_pdf import build_statistics_pdf, build_transfer_pdf
from verify_news import verify


def test_month_only_title_uses_source_file_year(monkeypatch, tmp_path: Path) -> None:
    source = tmp_path / "management.pdf"
    source.write_bytes(b"%PDF-placeholder")
    stamp = datetime(2026, 8, 15, 12, 0).timestamp()
    os.utime(source, (stamp, stamp))

    class Page:
        @staticmethod
        def extract_text() -> str:
            return "Nigeria Uyo Mission\nAugust Transfer"

    class FakePdf:
        pages = [Page()]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    monkeypatch.setattr(generate_news_format.pdfplumber, "open", lambda _path: FakePdf())
    assert generate_news_format.parse_transfer_title(source) == "AUGUST 2026 TRANSFER NEWS"


def test_zone_order_reads_the_index_even_without_spaces(monkeypatch, tmp_path: Path) -> None:
    lines = [
        "COMPLETE ZONE-BY-ZONE TRANSFER ASSIGNMENTS",
        "COMPLETEZONE-BY-ZONETRANSFERASSIGNMENTS",
        "01 UYOCENTRAL ZONE (15) 02 IBESIKPO ZONE (13)",
        "03 UYOZONE (12) 04 ITAM-NKEMBAZONE (12)",
        "# MISSIONARY NMEN ZONE NEW/EXISTING AREA COMPANION(S)",
        "02 IBESIKPO ZONE — 13 MISSIONARIES",
        "AP ASSISTANT TO ZL ZONE LEADER STL SISTER",
    ]

    class Page:
        @staticmethod
        def extract_text() -> str:
            return "\n".join(lines)

    class FakePdf:
        pages = [Page()]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

    monkeypatch.setattr(generate_news_format.pdfplumber, "open", lambda _path: FakePdf())
    assert generate_news_format.parse_zone_sections(tmp_path / "news.pdf") == [
        "UYO CENTRAL", "IBESIKPO", "UYO", "ITAM-NKEMBA",
    ]


def test_initial_suffix_names_share_a_base_name() -> None:
    assert generate_news_format._base_name("SMITH S.") == "SMITH"
    assert generate_news_format._base_name("SMITH.J") == "SMITH"
    assert generate_news_format._base_name("SMITH-JONES") == "SMITH-JONES"


def test_unmatched_news_row_is_blocking(tmp_path: Path) -> None:
    news = pd.DataFrame(
        [["ELDER MISSING", "JC", "UYO", "CENTRAL", "ELDER KNOWN", "SAME", "UYO", "CENTRAL", "JC"]],
        columns=[
            "Name of Missionary", "Assignment", "New/Existing Zone", "New/Existing Area",
            "New/Existing Companion(s)", "Status", "Previous Zone", "Previous Area", "Previous Assignment",
        ],
    )
    report = tmp_path / "current.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Zone", "District", "Area", "Missionary Name", "Type", "Position", "Companion"])
    sheet.append(["Uyo Zone", "Central", "Central", "Known, Elder", "Elder", "JC", "Missing, Elder (JC)"])
    workbook.save(report)

    result = verify(news, report)
    assert len(result.blocking_errors) == 1
    assert "cannot safely verify" in result.blocking_errors[0].detail


def _page_chars(words: list[tuple[str, float, float, float]]) -> list[dict]:
    """Synthetic pdfplumber chars for (text, x0, top, size) words."""
    chars = []
    for text, x0, top, size in words:
        for i, ch in enumerate(text):
            chars.append({"text": ch, "x0": x0 + 4 * i, "x1": x0 + 4 * (i + 1), "top": top, "size": size})
    return chars


def test_zone_page_blocks_are_found_when_the_second_block_moves_up() -> None:
    # Layout of a real export: the first district has no T-marker row, so
    # the second district prints ~11pt higher than the fixed TM_BLOCKS bands.
    chars = _page_chars([
        ("North Zone", 33, 36, 8.9),
        ("Harbor", 33, 65, 8.2),
        ("Lakeside", 37, 89, 6.0), ("Harbor", 105, 89, 6.0),
        ("ZL1", 49, 124, 6.0), ("DL", 119, 124, 6.0),
        ("Alpha", 35, 134, 7.4), ("Bravo", 66, 134, 7.4), ("Charlie", 104, 134, 7.4), ("Delta", 134, 134, 7.4),
        ("Hillside", 33, 173, 8.2),
        ("Hill Road", 37, 197, 6.0), ("MILL", 105, 197, 6.0),
        ("T", 59, 210, 5.2),
        ("SC", 50, 232, 6.0),
        ("Echo", 35, 242, 7.4), ("Foxtrot", 66, 242, 7.4), ("Golf", 104, 242, 7.4), ("Hotel", 134, 242, 7.4),
    ])
    assert generate_news_format._parse_block(chars, generate_news_format.TM_BLOCKS[1], "NORTH") == []

    blocks = generate_news_format._find_blocks(chars)
    areas = [a for block in blocks for a in generate_news_format._parse_block(chars, block, "NORTH")]
    assert [(a.district, a.area, [m.name for m in a.members]) for a in areas] == [
        ("Harbor", "Lakeside", ["Alpha", "Bravo"]),
        ("Harbor", "Harbor", ["Charlie", "Delta"]),
        ("Hillside", "Hill Road", ["Echo", "Foxtrot"]),
        ("Hillside", "MILL", ["Golf", "Hotel"]),
    ]
    assert areas[0].members[0].badge == "ZL1"
    assert areas[2].members[0].badge == "SC"
    assert areas[2].members[1].has_t


def test_roster_missionary_without_a_news_row_is_blocking(tmp_path: Path) -> None:
    columns = [
        "Name of Missionary", "Assignment", "New/Existing Zone", "New/Existing Area",
        "New/Existing Companion(s)", "Status", "Previous Zone", "Previous Area", "Previous Assignment",
    ]
    news = pd.DataFrame(
        [
            ["ELDER KNOWN", "SC", "UYO", "CENTRAL", "ELDER PARTNER", "SAME", "UYO", "CENTRAL", "SC"],
            ["ELDER PARTNER", "JC", "UYO", "CENTRAL", "ELDER KNOWN", "SAME", "UYO", "CENTRAL", "JC"],
        ],
        columns=columns,
    )
    rows = [
        ["Uyo Zone", "Central", "Central", "Known, Elder", "Elder", "(SC)", "Partner, Elder (JC)"],
        ["Uyo Zone", "Central", "Central", "Partner, Elder", "Elder", "(JC)", "Known, Elder (SC)"],
        ["Uyo Zone", "East", "Mill Road 1&amp;2/", "Skipped, Elder", "Elder", "(SC)", ""],
    ]
    reports = []
    for name in ("current.xlsx", "old.xlsx"):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Zone", "District", "Area", "Missionary Name", "Type", "Position", "Companion"])
        for row in rows:
            sheet.append(row)
        workbook.save(tmp_path / name)
        reports.append(tmp_path / name)

    result = verify(news, *reports)
    assert [e.row_name for e in result.blocking_errors] == ["ELDER SKIPPED"]
    assert "(Uyo / Mill Road 1&2)" in result.blocking_errors[0].detail
    assert "no row" in result.blocking_errors[0].detail


def _prev_record(assignment: str, companion: str = "") -> generate_news_format.MissionaryRecord:
    return generate_news_format.MissionaryRecord(
        display_name="ELDER X", title="ELDER", last_name="X", assignment=assignment,
        section_zone="UYO", assigned_zone="UYO", area="CENTRAL", companion=companion,
    )


def test_sa_is_not_inferred_from_a_previous_transfer() -> None:
    TM = generate_news_format.TMMissionary
    former_sa_pair = generate_news_format.AreaRecord(zone="UYO", district="CENTRAL", area="CENTRAL", members=[
        TM(name="Alpha", x=0, badge="SC", prev=_prev_record("SA", "ELDER BRAVO"), state="matched"),
        TM(name="Bravo", x=30, badge="JC", prev=_prev_record("SA", "ELDER ALPHA"), state="matched"),
    ])
    assert generate_news_format._group_assignments(former_sa_pair) == ["SC", "JC"]

    # An SA transferred away from the former SA companionship follows the
    # current DL/JC pairing rather than keeping the old special assignment.
    pair = generate_news_format.AreaRecord(zone="UYO", district="CENTRAL", area="CENTRAL", members=[
        TM(name="Delta", x=0, badge="DL", prev=_prev_record("DL"), state="matched"),
        TM(name="Echo", x=30, badge="JC", prev=_prev_record("SA", "ELDER FORMER COMPANION"), state="matched"),
    ])
    assert generate_news_format._group_assignments(pair) == ["DL", "JC"]


def test_manual_assignment_correction_accepts_a_full_missionary_name(tmp_path: Path) -> None:
    import verify_news

    corrections_file = tmp_path / "role-overrides.csv"
    corrections_file.write_text(
        "Name,Field,Value\nElder Erondu,Assignment,DL\nElder Asamoah,Assignment,SA\n",
        encoding="utf-8",
    )
    news = pd.DataFrame(
        [
            {"Name of Missionary": "ELDER ERONDU", "Assignment": "SA"},
            {"Name of Missionary": "ELDER ASAMOAH", "Assignment": "JC"},
        ]
    )
    corrections = []

    verify_news._apply_manual_corrections(news, corrections_file, corrections)

    assert news.at[0, "Assignment"] == "DL"
    assert news.at[1, "Assignment"] == "SA"


def test_report_sc_and_jc_do_not_overwrite_special_assignments(tmp_path: Path) -> None:
    columns = [
        "Name of Missionary", "Assignment", "New/Existing Zone", "New/Existing Area",
        "New/Existing Companion(s)", "Status", "Previous Zone", "Previous Area", "Previous Assignment",
    ]
    news = pd.DataFrame(
        [
            ["ELDER ALPHA", "SA", "UYO", "CENTRAL", "ELDER BRAVO", "STAYING", "UYO", "CENTRAL", "SA"],
            ["ELDER BRAVO", "SA", "UYO", "CENTRAL", "ELDER ALPHA", "STAYING", "UYO", "CENTRAL", "SA"],
        ],
        columns=columns,
    )
    rows = [
        ["Uyo Zone", "Central", "Central", "Alpha, Elder", "Elder", "(SC)", "Bravo, Elder (JC)"],
        ["Uyo Zone", "Central", "Central", "Bravo, Elder", "Elder", "(JC)", "Alpha, Elder (SC)"],
    ]
    reports = []
    for name in ("current.xlsx", "old.xlsx"):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["Zone", "District", "Area", "Missionary Name", "Type", "Position", "Companion"])
        for row in rows:
            sheet.append(row)
        workbook.save(tmp_path / name)
        reports.append(tmp_path / name)

    result = verify(news, *reports, manual_corrections_path=tmp_path / "none.csv")
    assert list(result.news["Assignment"]) == ["SA", "SA"]
    assert list(result.news["Previous Assignment"]) == ["SA", "SA"]
    assert list(result.news["Status"]) == ["STAYING", "STAYING"]
    assert result.corrections == []


def test_verified_zone_changes_keep_the_carried_over_zone_order() -> None:
    from data.generate_transfer_sheet import _restore_zone_order

    news = pd.DataFrame(
        [
            # Parser grouped BRAVO with IBESIKPO; verification corrected it.
            {"Name of Missionary": "ELDER ALPHA", "Previous Zone": "IBESIKPO", "New/Existing Zone": "IBESIKPO",
             "_zone_order": 0, "_prev_tier": 3, "_row_order": 5},
            {"Name of Missionary": "ELDER BRAVO", "Previous Zone": "UKAT-NSIT", "New/Existing Zone": "IBESIKPO",
             "_zone_order": 0, "_prev_tier": 3, "_row_order": 10 ** 9},
            {"Name of Missionary": "ELDER CHARLIE", "Previous Zone": "UYO", "New/Existing Zone": "UYO",
             "_zone_order": 1, "_prev_tier": 3, "_row_order": 7},
            {"Name of Missionary": "ELDER DELTA", "Previous Zone": "UKAT-NSIT", "New/Existing Zone": "UKAT-NSIT",
             "_zone_order": 2, "_prev_tier": 1, "_row_order": 9},
        ]
    )
    ordered = _restore_zone_order(news, ["IBESIKPO", "UYO", "UKAT-NSIT"])
    assert list(ordered["Name of Missionary"]) == ["ELDER ALPHA", "ELDER CHARLIE", "ELDER DELTA", "ELDER BRAVO"]


def test_current_leadership_roles_print_first_and_keep_pairs_together() -> None:
    from data.generate_transfer_sheet import _restore_zone_order

    news = pd.DataFrame(
        [
            {"Name of Missionary": "ELDER AP2", "Assignment": "AP2", "Previous Zone": "UYO", "New/Existing Zone": "UYO", "_row_order": 1},
            {"Name of Missionary": "ELDER SC", "Assignment": "SC", "Previous Zone": "UYO", "New/Existing Zone": "UYO", "_row_order": 2},
            {"Name of Missionary": "ELDER AP1", "Assignment": "AP1", "Previous Zone": "UYO", "New/Existing Zone": "UYO", "_row_order": 9},
            {"Name of Missionary": "ELDER ZL1", "Assignment": "ZL1", "Previous Zone": "UYO", "New/Existing Zone": "UYO", "_row_order": 3},
        ]
    )

    ordered = _restore_zone_order(news, ["UYO"])
    assert list(ordered["Name of Missionary"]) == ["ELDER AP1", "ELDER AP2", "ELDER ZL1", "ELDER SC"]


def test_stats_count_a_trio_as_one_companionship(sample_news: pd.DataFrame) -> None:
    transfer = pd.DataFrame(
        [{"Zone": "UYO", "District": "CENTRAL", "Area": "CENTRAL"}]
    )
    stats = compute_stats(sample_news, transfer, new_missionary_count=1)
    assert stats["Total Missionaries in Mission"] == 4
    assert stats["Companionships"] == 1
    assert stats["Zones"] == 1


def test_native_pdf_and_docx_outputs_are_structurally_valid(sample_workbook: Path, tmp_path: Path) -> None:
    news_pdf = build_transfer_pdf(sample_workbook, tmp_path / "news.pdf")
    stats_pdf = build_statistics_pdf(sample_workbook, tmp_path / "stats.pdf")
    news_docx = build_transfer_docx(sample_workbook, tmp_path / "news.docx", None)

    reader = PdfReader(news_pdf)
    assert len(reader.pages) >= 3
    pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "AUGUST 2026 TRANSFER NEWS" in pdf_text
    assert "ELDER ALPHA" in pdf_text
    assert len(PdfReader(stats_pdf).pages) == 1

    document = openpyxl.load_workbook(sample_workbook, read_only=True)
    assert {"News Format", "Transfer Sheet", "Meta"}.issubset(document.sheetnames)
    document.close()
    assert news_docx.exists() and news_docx.stat().st_size > 10_000
