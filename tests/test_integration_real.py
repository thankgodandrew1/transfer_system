from __future__ import annotations

import json
import os
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from docx.enum.section import WD_ORIENT

import app as webapp


@pytest.mark.skipif(os.getenv("RUN_REAL_DATA_TESTS") != "1", reason="local confidential fixtures are opt-in")
def test_real_inputs_generate_full_web_package(monkeypatch, tmp_path: Path) -> None:
    data_dir = Path(__file__).resolve().parents[1] / "data"
    inputs = {
        "management_pdf": data_dir / "current_transfer_management.pdf",
        "previous_news_pdf": data_dir / "old_transfer_news.pdf",
        "current_report": data_dir / "current_transfer_report.xlsx",
        "old_report": data_dir / "old_transfer_report.xlsx",
    }
    if not all(path.exists() for path in inputs.values()):
        pytest.skip("local transfer fixtures are not present")

    job_root = tmp_path / "jobs"
    job_root.mkdir()
    monkeypatch.setattr(webapp, "JOB_ROOT", job_root)
    monkeypatch.setattr(webapp, "ACCESS_KEY", "")
    webapp.app.config.update(TESTING=True, SECRET_KEY="integration-secret", SESSION_COOKIE_SECURE=False)
    client = webapp.app.test_client()
    client.get("/")
    with client.session_transaction() as session:
        token = session["csrf_token"]

    handles = {field: path.open("rb") for field, path in inputs.items()}
    try:
        response = client.post(
            "/generate",
            data={
                "csrf_token": token,
                "include_roster": "on",
                "include_stats": "on",
                **{field: (stream, inputs[field].name) for field, stream in handles.items()},
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
    finally:
        for stream in handles.values():
            stream.close()

    assert response.status_code == 200
    manifests = list(job_root.glob("*/manifest.json"))
    assert len(manifests) == 1
    manifest = json.loads(manifests[0].read_text(encoding="utf-8"))
    assert manifest["status"] == "complete", manifest.get("error")
    assert manifest["summary"]["new_missionaries"] == 16
    assert manifest["summary"]["blocking"] == 0
    kinds = {item["kind"] for item in manifest["files"]}
    assert {"zip", "xlsx", "csv", "log", "docx", "pdf"}.issubset(kinds)
    for item in manifest["files"]:
        assert (manifests[0].parent / item["relative_path"]).stat().st_size > 0

    review_workbook = next(
        manifests[0].parent / item["relative_path"]
        for item in manifest["files"]
        if item["label"] == "Review workbook"
    )
    workbook = openpyxl.load_workbook(review_workbook, read_only=False, data_only=True)
    news_sheet = workbook["News Format"]
    assert news_sheet.freeze_panes == "A2"
    assert news_sheet.sheet_view.showGridLines is False
    assert news_sheet.auto_filter.ref
    assert news_sheet["A1"].fill.fgColor.rgb.endswith("1A5C38")
    workbook.close()

    transfer_docx = next(
        manifests[0].parent / item["relative_path"]
        for item in manifest["files"]
        if item["label"] == "Transfer News - Word"
    )
    document = Document(transfer_docx)
    assert document.sections[0].orientation == WD_ORIENT.LANDSCAPE
    assert len(document.tables) >= 12
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "AUGUST 2026 TRANSFER NEWS" in document_text
    assert "{{" not in document_text


@pytest.mark.skipif(os.getenv("RUN_REAL_DATA_TESTS") != "1", reason="local confidential fixtures are opt-in")
def test_real_transfer_news_generates_sample_matched_movement(monkeypatch, tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    inputs = {
        "movement_previous_pdf": root / "data" / "old_transfer_news.pdf",
        "movement_current_pdf": Path(
            r"C:\Users\369659\Desktop\AP'S SPICY\TRANSFER\AUGUST 2026\Transfer_News_August_2026.pdf"
        ),
        "apartment_directory": Path(r"C:\Users\369659\Downloads\nigeria_uyo_apartment_map.json"),
    }
    if not all(path.exists() for path in inputs.values()):
        pytest.skip("local movement fixtures are not present")

    job_root = tmp_path / "movement-jobs"
    job_root.mkdir()
    monkeypatch.setattr(webapp, "JOB_ROOT", job_root)
    monkeypatch.setattr(webapp, "ACCESS_KEY", "")
    webapp.app.config.update(TESTING=True, SECRET_KEY="movement-integration-secret", SESSION_COOKIE_SECURE=False)
    client = webapp.app.test_client()
    client.get("/movement")
    with client.session_transaction() as session:
        token = session["csrf_token"]

    handles = {field: path.open("rb") for field, path in inputs.items()}
    try:
        response = client.post(
            "/movement/generate",
            data={
                "csrf_token": token,
                "manual_overrides": "Elder Twum: Anua Obio -> Mbierebe",
                **{field: (stream, inputs[field].name) for field, stream in handles.items()},
            },
            content_type="multipart/form-data",
            follow_redirects=True,
        )
    finally:
        for stream in handles.values():
            stream.close()

    assert response.status_code == 200
    manifest_path = next(job_root.glob("*/manifest.json"))
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert manifest["status"] == "complete", manifest.get("error")
    assert manifest["kind"] == "movement"
    assert manifest["summary"]["movements"] == 70
    assert manifest["summary"]["zones"] == 12
    assert manifest["summary"]["new_arrivals"] == 18
    assert manifest["summary"]["released"] == 6

    movement_docx = next(
        manifest_path.parent / item["relative_path"]
        for item in manifest["files"]
        if item["label"] == "Transfer Movement - Word"
    )
    document = Document(movement_docx)
    movement_tables = [
        table for table in document.tables
        if table.rows and "MISSIONARY" in "|".join(cell.text for cell in table.rows[0].cells).upper()
    ]
    assert len(movement_tables) == 12
    assert sum(len(table.rows) - 1 for table in movement_tables) == 70
    assert all(row.cells[4].text == "" for table in movement_tables for row in table.rows[1:])
