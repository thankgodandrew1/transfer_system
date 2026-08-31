from __future__ import annotations

import io
import json
from pathlib import Path

from docx import Document

import app as webapp
import movement_plan
from movement_plan import (
    ApartmentDirectory,
    Assignment,
    DirectoryEntry,
    build_movement_docx,
    compare_assignments,
    group_movements,
    load_apartment_directory,
    parse_manual_overrides,
    parse_transfer_text,
)


def _directory() -> ApartmentDirectory:
    return ApartmentDirectory(
        [
            DirectoryEntry("OLD ZONE", "OLD APARTMENT", "OLD AREA"),
            DirectoryEntry("NEW ZONE", "NEW APARTMENT", "NEW AREA"),
            DirectoryEntry("OLD ZONE", "OLD APARTMENT", "UNCHANGED AREA"),
        ]
    )


def test_compare_skips_arrivals_and_releases_and_groups_by_previous_zone() -> None:
    previous = [
        Assignment("Elder Alpha", "OLD ZONE", "OLD AREA", 1),
        Assignment("Sister Same", "OLD ZONE", "UNCHANGED AREA", 2),
        Assignment("Elder Released", "OLD ZONE", "OLD AREA", 3),
    ]
    current = [
        Assignment("ELDERALPHA", "NEW ZONE", "NEW AREA", 1),
        Assignment("Sister Same", "OLD ZONE", "UNCHANGED AREA", 2),
        Assignment("Sister New", "NEW ZONE", "NEW AREA", 3),
    ]
    result = compare_assignments(previous, current, _directory(), default_status="scheduled")

    assert len(result.movements) == 1
    movement = result.movements[0]
    assert movement.missionary == "ELDERALPHA"
    assert movement.previous_zone == "OLD ZONE"
    assert movement.from_apartment == "OLD APARTMENT"
    assert movement.to_apartment == "NEW APARTMENT"
    assert movement.status == ""
    assert result.released_total == 1
    assert result.new_arrivals_total == 1
    assert result.unchanged_total == 1
    assert group_movements(result.movements)[0][0] == "OLD ZONE"


def test_manual_override_supports_status_column() -> None:
    overrides = parse_manual_overrides(
        "Elder Alpha: Historic Apt -> Correct Apt | CONFIRMED\n"
        "Sister Same | Apt One | Apt Two | READY",
        "PENDING",
    )
    assert [(item.name, item.status) for item in overrides] == [
        ("Elder Alpha", "CONFIRMED"),
        ("Sister Same", "READY"),
    ]


def test_plain_text_fallback_reads_delimited_table() -> None:
    parsed = parse_transfer_text(
        "Missionary | Zone | Area\nElder Alpha | Old Zone | Old Area\nSister Beta | New Zone | New Area"
    )
    assert [item.name for item in parsed.assignments] == ["Elder Alpha", "Sister Beta"]


def test_pdf_parser_rejoins_rows_split_across_pages(monkeypatch) -> None:
    header = ["#", "MISSIONARY", "ASSIGNMENT", "ZONE", "NEW/EXISTING AREA", "COMPANION(S)"]

    class FakePage:
        def __init__(self, table, text=""):
            self._table = table
            self._text = text

        def extract_text(self, **_kwargs):
            return self._text

        def extract_tables(self):
            return [self._table]

    class FakePdf:
        pages = [
            FakePage(
                [header, ["7", "SISTER", "SC", "IKOT", "IKOT IKWA", "SISTER BOATENG"]],
                "AUGUST 2026 TRANSFER NEWS",
            ),
            FakePage(
                [
                    header,
                    ["", "OKOMO", "", "USEKONG", "", ""],
                    ["1\n0", "ELDER ALPHA", "JC", "UYO", "OLD AREA", "ELDER BETA"],
                ]
            ),
        ]

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

    monkeypatch.setattr(movement_plan.pdfplumber, "open", lambda _source: FakePdf())
    parsed = movement_plan.parse_transfer_pdf("split-across-pages.pdf")

    assert [(item.name, item.zone, item.area) for item in parsed.assignments] == [
        ("SISTER OKOMO", "IKOT USEKONG", "IKOT IKWA"),
        ("ELDER ALPHA", "UYO", "OLD AREA"),
    ]


def test_json_directory_is_loaded_and_exported_as_excel_friendly_csv(tmp_path: Path) -> None:
    source = tmp_path / "apartments.json"
    source.write_text(
        json.dumps(
            {
                "area_to_apartment": {
                    "OLD AREA": {"apartment": "OLD APARTMENT", "zone": "OLD ZONE"},
                    "NEW AREA": {"apartment": "NEW APARTMENT", "zone": "NEW ZONE"},
                }
            }
        ),
        encoding="utf-8",
    )
    directory = load_apartment_directory(source)
    assert directory.find("oldarea").apartment == "OLD APARTMENT"
    exported = directory.to_csv_bytes().decode("utf-8-sig")
    assert exported.startswith("Zone,Apartment,Area")
    assert "OLD ZONE,OLD APARTMENT,OLD AREA" in exported


def test_movement_docx_has_requested_status_column(tmp_path: Path) -> None:
    comparison = compare_assignments(
        [Assignment("Elder Alpha", "OLD ZONE", "OLD AREA", 1)],
        [Assignment("Elder Alpha", "NEW ZONE", "NEW AREA", 1)],
        _directory(),
    )
    output = build_movement_docx(
        comparison.movements,
        tmp_path / "movement.docx",
        mission_name="Nigeria Uyo Mission",
        previous_title="June / July 2026 Transfer News",
        current_title="August 2026 Transfer News",
        president="President Sample",
        prepared_by="Mission Office",
    )
    document = Document(output)
    movement_tables = [
        table for table in document.tables
        if table.rows and "MISSIONARY" in "|".join(cell.text for cell in table.rows[0].cells).upper()
    ]
    assert len(movement_tables) == 1
    assert [cell.text for cell in movement_tables[0].rows[0].cells] == [
        "#",
        "MISSIONARY",
        "FROM APARTMENT",
        "TO APARTMENT",
        "STATUS",
    ]
    assert movement_tables[0].cell(1, 4).text == ""


def _client(monkeypatch, tmp_path: Path):
    job_root = tmp_path / "jobs"
    job_root.mkdir()
    monkeypatch.setattr(webapp, "JOB_ROOT", job_root)
    monkeypatch.setattr(webapp, "ACCESS_KEY", "")
    webapp.app.config.update(TESTING=True, SECRET_KEY="movement-secret", SESSION_COOKIE_SECURE=False)
    return webapp.app.test_client()


def test_movement_page_and_completed_job(monkeypatch, tmp_path: Path) -> None:
    client = _client(monkeypatch, tmp_path)
    page = client.get("/movement")
    assert page.status_code == 200
    assert b"Generate Movement Plan" in page.data
    assert b"STATUS" in page.data
    with client.session_transaction() as session:
        token = session["csrf_token"]

    def fake_run(job_dir: Path, _values):
        output = job_dir / "outputs" / "movement.docx"
        output.parent.mkdir(parents=True)
        output.write_bytes(b"docx-test")
        return ([{
            "name": output.name,
            "label": "Transfer Movement - Word",
            "kind": "docx",
            "size": output.stat().st_size,
            "relative_path": "outputs/movement.docx",
        }], {
            "movements": 2,
            "zones": 1,
            "matched": 3,
            "released": 1,
            "new_arrivals": 1,
            "unchanged": 1,
            "previous_total": 4,
            "current_total": 4,
            "manual_overrides": 1,
        }, [])

    monkeypatch.setattr(webapp, "_run_movement_generation", fake_run)
    response = client.post(
        "/movement/generate",
        data={"csrf_token": token, "apartment_directory": (io.BytesIO(b"ignored"), "directory.csv")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Movement package is ready" in response.data
    assert b"New arrivals skipped" in response.data


def test_apartment_directory_template_download(monkeypatch, tmp_path: Path) -> None:
    client = _client(monkeypatch, tmp_path)
    response = client.get("/movement/apartment-directory-template.csv")
    assert response.status_code == 200
    assert response.data.startswith(b"Zone,Apartment,Area")
    assert "attachment" in response.headers["Content-Disposition"]
